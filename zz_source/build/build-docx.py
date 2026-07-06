#!/usr/bin/env python3
"""
Convert the Markdown resources into editable Word (.docx) files.

Builds clean DOCX directly with python-docx (no LibreOffice/pandoc needed),
handling headings, bold/italic/inline-code, tables, bullet/numbered lists,
fenced code blocks, blockquotes, horizontal rules and <details> answer
sections (rendered open so answers are visible/editable).

Usage:
    python3 revision-tools/build-docx.py                     # build the default set
    python3 revision-tools/build-docx.py <srcdir> <outdir>   # one folder
"""
import sys, os, re, glob

try:
    import docx
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx not installed. Run: pip install python-docx")

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# NOTE: no `__bold__` variant — answer blanks like "Name: ______" must stay literal.
INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+?`)')
ESC_STAR = '\x00'   # placeholder for the \* escape so it never pairs as emphasis


def _emit(paragraph, text, bold=False, italic=False, code=False):
    r = paragraph.add_run(text.replace(ESC_STAR, '*'))
    r.bold = bold; r.italic = italic
    if code:
        r.font.name = 'Consolas'; r.font.size = Pt(9.5)


def _parse_inline(paragraph, text, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            inner = part[2:-2]
            # allow `code` nested inside bold
            for sub in re.split(r'(`[^`]+?`)', inner):
                if not sub:
                    continue
                if sub.startswith('`') and sub.endswith('`') and len(sub) > 2:
                    _emit(paragraph, sub[1:-1], bold=True, italic=italic, code=True)
                else:
                    _emit(paragraph, sub, bold=True, italic=italic)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            _emit(paragraph, part[1:-1], bold=bold, italic=italic, code=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            _emit(paragraph, part[1:-1], bold=bold, italic=True)
        else:
            _emit(paragraph, part, bold=bold, italic=italic)


def add_runs(paragraph, text):
    """Add text to a paragraph, parsing **bold**, *italic* and `code`
    (with `code` allowed inside bold, and \\* kept as a literal star)."""
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)   # links -> just the label
    text = text.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
    text = text.replace('\\*', ESC_STAR)
    _parse_inline(paragraph, text)


def shade(cell_or_para, fill):
    el = cell_or_para._tc if hasattr(cell_or_para, '_tc') else cell_or_para._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_para, '_tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill)
    pr.append(shd)


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        convert_text(f.read(), docx_path)


def convert_text(text, docx_path):
    # force any collapsible answers open, drop the html tags
    text = re.sub(r'</?details[^>]*>', '', text)
    text = text.replace('<summary>', '**').replace('</summary>', '**')
    lines = text.split('\n')

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]

        # fenced code block
        if line.lstrip().startswith('```'):
            i += 1
            code = []
            while i < n and not lines[i].lstrip().startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            shade(p, 'F2F2F2')
            r = p.add_run('\n'.join(code))
            r.font.name = 'Consolas'; r.font.size = Pt(9)
            continue

        # table (header row followed by |---| separator)
        if line.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lines[i+1]) and '-' in lines[i+1]:
            block = []
            while i < n and lines[i].strip().startswith('|'):
                block.append(lines[i]); i += 1
            def cells(row):
                row = row.strip()
                if row.startswith('|'): row = row[1:]
                if row.endswith('|'): row = row[:-1]
                return [c.strip() for c in row.split('|')]
            header = cells(block[0])
            data = [cells(r) for r in block[2:]]
            tbl = doc.add_table(rows=1, cols=len(header)); tbl.style = 'Table Grid'
            for j, h in enumerate(header):
                c = tbl.rows[0].cells[j]; c.paragraphs[0].text = ''
                add_runs(c.paragraphs[0], h)
                for run in c.paragraphs[0].runs: run.bold = True
                shade(c, 'D9E2F3')
            for drow in data:
                cs = tbl.add_row().cells
                for j in range(len(header)):
                    cs[j].paragraphs[0].text = ''
                    add_runs(cs[j].paragraphs[0], drow[j] if j < len(drow) else '')
            doc.add_paragraph()
            continue

        # headings
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=min(level, 4))
            add_runs(h, m.group(2))
            i += 1
            continue

        # horizontal rule
        if re.match(r'^\s*([-*_])\1{2,}\s*$', line):
            p = doc.add_paragraph(); pr = p._p.get_or_add_pPr()
            pb = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'AAAAAA')
            pb.append(bottom); pr.append(pb)
            i += 1
            continue

        # blockquote
        if line.lstrip().startswith('>'):
            content = re.sub(r'^\s*>\s?', '', line)
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            shade(p, 'FFF8E7')
            add_runs(p, content)
            i += 1
            continue

        # bullet list
        m = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(2)); i += 1
            continue
        # numbered list
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2)); i += 1
            continue

        # answer-writing space sentinel (@@SPACE:n@@) -> typeable answer box:
        # a single-cell bordered table with a minimum height that grows as
        # the student types (exam-paper style answer box).
        ms = re.match(r'^@@SPACE:(\d+)@@$', line.strip())
        if ms:
            n_lines = int(ms.group(1))
            if n_lines == 0:               # answer lives in the table/code above
                i += 1
                continue
            tbl = doc.add_table(rows=1, cols=1)
            tbl.style = 'Table Grid'
            row = tbl.rows[0]
            tr_pr = row._tr.get_or_add_trPr()
            h = OxmlElement('w:trHeight')
            h.set(qn('w:val'), str(n_lines * 340))        # ~0.24in per answer line
            h.set(qn('w:hRule'), 'atLeast')               # grows with typed content
            tr_pr.append(h)
            cell = row.cells[0]
            cp0 = cell.paragraphs[0]
            cp0.paragraph_format.space_before = Pt(2)
            doc.add_paragraph()                            # spacer after the box
            i += 1
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)


QA_FOLDERS = {'subtopic-quizzes', 'worksheets', 'mini-papers', 'mock-papers',
              'homework', 'recap-checkpoints'}
ANSWER_HEADING = re.compile(r'^\s*#{2,3}\s+(answer key|answers|mark scheme)', re.I)
MARK_TAG = re.compile(r'\[(\d+)\]|\((\d+)\s*marks?\)')


def split_qa(text):
    lines = text.split('\n')
    for idx, l in enumerate(lines):
        if ANSWER_HEADING.match(l):
            return '\n'.join(lines[:idx]).rstrip(), '\n'.join(lines[idx:]).strip()
    return text, None


SPACE_SENTINEL = re.compile(r'^@@SPACE:(\d+)@@$')


def inject_space(text):
    """Add a typeable answer box after each question line carrying a mark tag.
    The box goes after any code fence or table attached to the stem, never
    inside it. An explicit @@SPACE:n@@ sentinel directly after the question
    block suppresses the automatic box (n=0 = answer written in the table /
    code gaps themselves)."""
    lines = text.split('\n')
    out, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.lstrip()
        if stripped.startswith('```'):
            out.append(line); i += 1
            while i < n:
                out.append(lines[i]); i += 1
                if lines[i - 1].lstrip().startswith('```'):
                    break
            continue
        out.append(line); i += 1
        if stripped.startswith('#') or stripped.startswith('|'):
            continue
        tags = MARK_TAG.findall(line)
        if not tags:
            continue
        # consume material attached to the stem: blank lines, fences, tables
        while i < n:
            s = lines[i].strip()
            if s == '':
                out.append(lines[i]); i += 1
            elif s.startswith('```'):
                out.append(lines[i]); i += 1
                while i < n:
                    out.append(lines[i]); i += 1
                    if lines[i - 1].lstrip().startswith('```'):
                        break
            elif s.startswith('|'):
                while i < n and lines[i].lstrip().startswith('|'):
                    out.append(lines[i]); i += 1
            else:
                break
        # a @@SPACE sentinel anywhere before the next question stem / heading
        # means the author controls this question's answer space
        j, manual = i, False
        while j < n:
            s2 = lines[j].strip()
            if s2.startswith('```'):
                j += 1
                while j < n and not lines[j].lstrip().startswith('```'):
                    j += 1
                j += 1
                continue
            if SPACE_SENTINEL.match(s2):
                manual = True
                break
            if s2.startswith('#') or (not s2.startswith('|') and MARK_TAG.search(s2)):
                break
            j += 1
        if manual:
            continue
        nums = [int(x) for pair in tags for x in pair if x]
        marks = max(nums) if nums else 2
        out.append(f'@@SPACE:{min(12, max(2, round(marks * 1.4)))}@@')
    return '\n'.join(out)


def _h1(text):
    for l in text.split('\n'):
        if l.startswith('# '):
            return l[2:].strip()
    return 'Worksheet'


def build_folder(srcdir, outdir):
    folder = os.path.basename(srcdir.rstrip('/'))
    split = folder in QA_FOLDERS
    mds = sorted(f for f in glob.glob(os.path.join(srcdir, '*.md'))
                 if os.path.basename(f).lower() != 'readme.md')
    made = 0
    for md in mds:
        base = os.path.splitext(os.path.basename(md))[0]
        text = open(md, encoding='utf-8').read()
        if split:
            questions, answers = split_qa(text)
            convert_text(inject_space(questions), os.path.join(outdir, base + '.docx'))
            made += 1
            if answers:
                a_doc = f"# {_h1(text)} — ANSWER SHEET\n\n*Separate answer sheet / mark scheme.*\n\n" + answers
                convert_text(a_doc, os.path.join(outdir, base + '-ANSWERS.docx'))
                made += 1
        else:
            convert(md, os.path.join(outdir, base + '.docx'))
            made += 1
    return made, [os.path.basename(m) for m in mds]


DEFAULT_JOBS = [
    ('zz_source/component-01-computer-systems',                 'Word-Documents/01-Computer-Systems'),
    ('zz_source/component-02-algorithms-and-programming',       'Word-Documents/02-Algorithms-and-Programming'),
    ('zz_source/component-03-04-programming-project',           'Word-Documents/03-04-Programming-Project'),
    ('zz_source/revision-tools/guides',                         'Word-Documents/Guides'),
    ('zz_source/revision-tools/knowledge-organisers',           'Word-Documents/Knowledge-Organisers'),
    ('zz_source/revision-tools/worksheets',                     'Word-Documents/Worksheets'),
    ('zz_source/revision-tools/revision-games',                 'Word-Documents/Revision-Games'),
    ('zz_source/revision-tools/subtopic-quizzes',               'Word-Documents/Subtopic-Quizzes'),
    ('zz_source/revision-tools/mini-papers',                    'Word-Documents/Mini-Papers'),
    ('zz_source/revision-tools/mock-papers',                    'Word-Documents/Mock-Papers'),
    ('zz_source/revision-tools/a-star-pack',                    'Word-Documents/A-Star-Pack'),
    ('zz_source/revision-tools/programming-workbook',           'Word-Documents/Programming-Workbook'),
    ('zz_source/revision-tools/practice-questions',             'Word-Documents/Practice-Questions'),
    ('zz_source/revision-tools/scheme-of-work',                 'Word-Documents/Scheme-of-Work'),
    ('zz_source/revision-tools/homework',                       'Word-Documents/Homework'),
    ('zz_source/revision-tools/recap-checkpoints',              'Word-Documents/Recap-Checkpoints'),
    ('zz_source/revision-tools/subtopic-revision',              'Word-Documents/Subtopic-Revision'),
    ('zz_source/revision-tools/nea-pack',                       'Word-Documents/NEA-Pack'),
    ('zz_source/revision-tools/course-guide',                   'Word-Documents/Course-Guide'),
    ('zz_source/revision-tools/lesson-activities',               'Word-Documents/Lesson-Activities'),
]


def main():
    if len(sys.argv) == 3:
        jobs = [(sys.argv[1], sys.argv[2])]
    else:
        jobs = [(os.path.join(REPO, s), os.path.join(REPO, o)) for s, o in DEFAULT_JOBS]
    total = 0
    for src, out in jobs:
        if not os.path.isdir(src):
            print(f"skip (missing): {src}"); continue
        made, names = build_folder(src, out)
        print(f"{os.path.relpath(src, REPO)} -> {os.path.relpath(out, REPO)}: {made} docx")
        total += made
    # README-only content files (NEA guide, practice-questions guide)
    if len(sys.argv) != 3:
        for s, o in [
            ("zz_source/component-03-04-programming-project/README.md",
             "Word-Documents/03-04-Programming-Project/Programming-Project-NEA-Guide.docx"),
            ("zz_source/revision-tools/practice-questions/README.md",
             "Word-Documents/Practice-Questions/Practice-Questions-Guide.docx"),
            ("zz_source/teacher-toolkit/medium-term-plan.md",
             "Teacher-Toolkit/Planning/medium-term-plan.docx"),
        ]:
            sp, op = os.path.join(REPO, s), os.path.join(REPO, o)
            if os.path.isfile(sp):
                convert(sp, op); total += 1
    print(f"\nDone. {total} DOCX generated.")


if __name__ == '__main__':
    main()
