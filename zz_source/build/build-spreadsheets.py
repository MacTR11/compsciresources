#!/usr/bin/env python3
"""
Build the planning/revision spreadsheets:
  - Flashcards.xlsx        one sheet per topic (Front | Back), generated from
                           the single tagged bank flashcards/by-subtopic.csv
  - Year-12-SOW.xlsx       the master weekly plan, generated from sow/year12.json
Also regenerates the per-topic flashcard CSVs (staging) from the tagged bank.

Output: zz_source/_staging/xl/  (and _staging/flashcards/ for the CSVs)
Usage:  python3 zz_source/build/build-spreadsheets.py
"""
import os, csv, json, re
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from taxonomy import REPO, TOPIC_TITLE, TOPIC_OF, SUB_STEM

OUT = os.path.join(REPO, "zz_source/_staging/xl")
FC_OUT = os.path.join(REPO, "zz_source/_staging/flashcards")
BANK = os.path.join(REPO, "zz_source", "revision-tools", "flashcards", "by-subtopic.csv")
SOW = os.path.join(REPO, "zz_source", "sow", "year12.json")

HEAD_FILL = PatternFill("solid", fgColor="2B6CB0")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="CCCCCC")] * 4)


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def load_bank():
    """by-subtopic.csv -> OrderedDict topic code ('1.1') -> [(q, a), ...]"""
    topics = OrderedDict()
    with open(BANK, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))[1:]
    for code, q, a in (r[:3] for r in rows if len(r) >= 3):
        topic = TOPIC_OF.get(code.strip())
        if topic:
            topics.setdefault(topic, []).append((q, a))
    return topics


def build_flashcards():
    topics = load_bank()
    # regenerate the per-topic CSVs (consumed by the Topic Library)
    os.makedirs(FC_OUT, exist_ok=True)
    for topic, cards in topics.items():
        with open(os.path.join(FC_OUT, f"{topic}.csv"), "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(cards)
    # the all-topics workbook
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for topic, cards in topics.items():
        ws = wb.create_sheet(title=topic)
        ws.append(["Front (question / term)", "Back (answer)"])
        for front, back in cards:
            ws.append([front, back])
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 70
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP; cell.border = THIN
        style_header(ws, 2)
    os.makedirs(OUT, exist_ok=True)
    wb.save(os.path.join(OUT, "Flashcards.xlsx"))
    print(f"  Flashcards.xlsx ({sum(len(c) for c in topics.values())} cards) + per-topic CSVs")


def _nice(stems):
    return ", ".join(s.split("-")[0] if s[0].isdigit() else s for s in stems)


def build_sow():
    weeks = json.load(open(SOW, encoding="utf-8"))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Year 12 Master SOW"
    ws.append(["Week", "Taught this week", "Class quiz (retrieval)", "Homework",
               "Consolidation", "Assessment", "Notes"])
    for w in weeks:
        hw = _nice(w["homework"]) or ""
        hw_full = (f"Homework {hw}: " if hw else "") + w["homework_note"]
        cons_bits = []
        if w["recap"]: cons_bits.append("Recap: " + _nice(w["recap"]))
        if w["minipaper"]: cons_bits.append("Mini-paper: " + _nice(w["minipaper"]))
        if w["worksheet"]: cons_bits.append("Worksheet: " + _nice(w["worksheet"]))
        cons_full = ("; ".join(cons_bits) + ". " if cons_bits else "") + w["consolidation_note"]
        ws.append([w["folder"], ", ".join(w["taught"]), _nice(w["class_quiz"]),
                   hw_full, cons_full, "; ".join(w["assessment"]), w["other_note"]])
    widths = {"A": 34, "B": 13, "C": 16, "D": 52, "E": 52, "F": 20, "G": 34}
    for col, wd in widths.items():
        ws.column_dimensions[col].width = wd
    style_header(ws, 7)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP; cell.border = THIN
        if not re.match(r"Week \d", str(row[0].value or "")):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="EFEFEF")
    os.makedirs(OUT, exist_ok=True)
    wb.save(os.path.join(OUT, "Year-12-SOW.xlsx"))
    print("  Year-12-SOW.xlsx")


if __name__ == "__main__":
    build_flashcards()
    build_sow()
    print("Done -> zz_source/_staging/xl/")
