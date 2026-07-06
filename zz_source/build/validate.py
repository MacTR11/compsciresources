#!/usr/bin/env python3
"""
Mechanical quality gate over the assembled tree. Run after assemble.py
(build.py does this automatically). Exits non-zero on any defect:

  - every .docx opens; no build sentinels (@@SPACE) or unrendered markdown
    (**, leading ##, |---, ``` fences) outside code-styled runs
  - every .pdf is non-trivial (> 5 KB)
  - every .pptx opens and has at least 5 slides
  - every .xlsx opens
  - INDEX.md references exactly the files that exist (1:1)
  - sources lint: no UPPERCASE pseudocode dialect in Component 2 sources
"""
import os, re, sys, glob
from urllib.parse import unquote

from taxonomy import REPO

TOP = ["0 Planning", "1 Year 12 Weekly", "2 Topic Library",
       "3 Assessments - TEACHER ONLY", "4 NEA and Programming",
       "5 Reference and Stretch"]

problems = []


def check_docx():
    import docx
    for top in TOP:
        for dirpath, _, files in os.walk(os.path.join(REPO, top)):
            for fn in files:
                if not fn.endswith(".docx"):
                    continue
                p = os.path.join(dirpath, fn)
                rel = os.path.relpath(p, REPO)
                try:
                    d = docx.Document(p)
                except Exception as e:
                    problems.append(f"docx will not open: {rel} ({e})")
                    continue
                paras = list(d.paragraphs)
                for t in d.tables:
                    for row in t.rows:
                        for c in row.cells:
                            paras.extend(c.paragraphs)
                if not any(p_.text.strip() for p_ in paras):
                    problems.append(f"docx is empty: {rel}")
                for p_ in paras:
                    txt = p_.text
                    if not txt.strip():
                        continue
                    is_code = any(r.font.name == "Consolas" for r in p_.runs)
                    if "@@SPACE" in txt:
                        problems.append(f"sentinel leaked: {rel}: {txt[:60]!r}")
                    if is_code:
                        continue
                    if "**" in txt:
                        problems.append(f"unrendered bold: {rel}: {txt[:60]!r}")
                    if re.match(r"^\s*#{1,6}\s", txt):
                        problems.append(f"unrendered heading: {rel}: {txt[:60]!r}")
                    if "|---" in txt or txt.strip().startswith("```"):
                        problems.append(f"unrendered md block: {rel}: {txt[:60]!r}")


def check_pdf_pptx_xlsx():
    for top in TOP:
        for dirpath, _, files in os.walk(os.path.join(REPO, top)):
            for fn in files:
                p = os.path.join(dirpath, fn)
                rel = os.path.relpath(p, REPO)
                if fn.endswith(".pdf") and os.path.getsize(p) < 5000:
                    problems.append(f"suspiciously small pdf: {rel} ({os.path.getsize(p)} bytes)")
                elif fn.endswith(".pptx"):
                    try:
                        from pptx import Presentation
                        if len(Presentation(p).slides) < 5:
                            problems.append(f"deck has under 5 slides: {rel}")
                    except Exception as e:
                        problems.append(f"pptx will not open: {rel} ({e})")
                elif fn.endswith(".xlsx"):
                    try:
                        import openpyxl
                        openpyxl.load_workbook(p, read_only=True).close()
                    except Exception as e:
                        problems.append(f"xlsx will not open: {rel} ({e})")


def check_index():
    idx = os.path.join(REPO, "INDEX.md")
    if not os.path.exists(idx):
        problems.append("INDEX.md missing")
        return
    linked = set()
    for m in re.finditer(r"\]\(([^)]+)\)", open(idx, encoding="utf-8").read()):
        linked.add(unquote(m.group(1)))
    actual = set()
    for top in TOP:
        for dirpath, _, files in os.walk(os.path.join(REPO, top)):
            for fn in files:
                actual.add(os.path.relpath(os.path.join(dirpath, fn), REPO))
    for f in sorted(actual - linked):
        problems.append(f"file not in INDEX.md: {f}")
    for f in sorted(linked - actual):
        problems.append(f"INDEX.md links to missing file: {f}")


def check_sources():
    for pat in ("zz_source/component-02-*/**/*.md", "zz_source/revision-tools/mini-papers/2.*.md",
                "zz_source/revision-tools/subtopic-revision/2.*.md"):
        for f in glob.glob(os.path.join(REPO, pat), recursive=True):
            text = open(f, encoding="utf-8").read()
            in_fence, lineno = False, 0
            for line in text.split("\n"):
                lineno += 1
                if line.lstrip().startswith("```"):
                    in_fence = not in_fence
                    continue
                if in_fence and re.search(r"\b(ENDWHILE|ENDFUNCTION|ENDIF|REPEAT|LEN\()", line):
                    problems.append(f"non-OCR pseudocode dialect: "
                                    f"{os.path.relpath(f, REPO)}:{lineno}: {line.strip()[:50]!r}")


def main():
    check_docx()
    check_pdf_pptx_xlsx()
    check_index()
    check_sources()
    if problems:
        print(f"VALIDATION FAILED - {len(problems)} problem(s):")
        for p in problems[:80]:
            print("  -", p)
        if len(problems) > 80:
            print(f"  ... and {len(problems) - 80} more")
        sys.exit(1)
    print("Validation passed: all documents open cleanly, INDEX is 1:1, sources lint clean.")


if __name__ == "__main__":
    main()
